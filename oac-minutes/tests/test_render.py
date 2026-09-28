from datetime import date

from docx import Document
from openpyxl import load_workbook

from oacminutes.carryforward import carry_forward
from oacminutes.formats import load_builtin_format
from oacminutes.models import Attendee, Firm, Meeting, MeetingExtraction, NewItem, NextMeeting, PriorItemUpdate, Project, SectionNotes
from oacminutes.render_common import output_basename, slugify, status_label
from oacminutes.render_docx import render_docx
from oacminutes.render_md import render_markdown
from oacminutes.render_xlsx import render_csv, render_xlsx


def scenario(format_key="standard_opm"):
    firm = Firm(name="Acme OPM", format=load_builtin_format(format_key))
    project = Project(firm_id=firm.id, name="Town Library", number="24-1", owner="Town of X", prepared_by="Pat PM")
    m1_ext = MeetingExtraction(new_items=[
        NewItem(description="Submit schedule", responsible="GC", due_date="2026-09-05", section_key="schedule"),
        NewItem(description="Answer RFI 4", responsible="Architect", due_date="2026-09-20", section_key="rfis_submittals"),
        NewItem(description="Pick colors", responsible="Owner", due_date="2026-09-30", section_key="owner_items"),
    ])
    r1 = carry_forward([], m1_ext, project_id=project.id, meeting_no=1, meeting_date=date(2026, 9, 1), numbering=firm.format.numbering)
    m2_ext = MeetingExtraction(
        attendees=[Attendee(name="Pat PM", company="Acme OPM", role="OPM"), Attendee(name="Sam", company="GC Co", role="GC", present=False)],
        sections=[SectionNotes(section_key="safety", bullets=["No incidents."]), SectionNotes(section_key="schedule", bullets=["Schedule is late.", "Steel on 10/19."])],
        decisions=["Approved PCO 1 at $100."],
        prior_item_updates=[
            PriorItemUpdate(item_number=r1.items[0].item_number, status="open", note="Still late."),
            PriorItemUpdate(item_number=r1.items[1].item_number, status="closed", note="Answered 9/7."),
        ],
        new_items=[NewItem(description="Order doors", responsible="GC", due_date="2026-09-15", section_key="rfis_submittals")],
        next_meeting=NextMeeting(date="2026-09-15", time="9:00 AM", location="Trailer"),
    )
    meeting = Meeting(project_id=project.id, meeting_no=2, meeting_date=date(2026, 9, 8), extraction=m2_ext, location="Trailer")
    r2 = carry_forward(r1.items, m2_ext, project_id=project.id, meeting_no=2, meeting_date=date(2026, 9, 8), numbering=firm.format.numbering)
    return firm, project, meeting, r2


def test_status_labels():
    firm, project, meeting, r2 = scenario()
    by = r2.by_number()
    labels = {n: status_label(i, firm.format, meeting_no=2, as_of=date(2026, 9, 8)) for n, i in by.items()}
    assert labels == {"1.01": "OVERDUE", "1.02": "Closed", "1.03": "Open", "2.01": "New"}


def test_markdown_contains_everything():
    firm, project, meeting, r2 = scenario()
    md = render_markdown(firm, project, meeting, r2)
    assert md.startswith("# OAC Meeting Minutes")
    assert "| **Project** | Town Library |" in md
    assert "## Attendees" in md and "| Sam | GC Co | GC | No |" in md
    assert "## Safety" in md and "- No incidents." in md
    assert "## Decisions" in md and "Approved PCO 1" in md
    assert "| 1.01 | Submit schedule | GC | 09/01/2026 | 09/05/2026 | OVERDUE | Still late. |" in md
    assert "| 1.02 | Answer RFI 4 | Architect | 09/01/2026 | 09/20/2026 | Closed | Answered 9/7. |" in md
    assert "| 1.03 | Pick colors | Owner | 09/01/2026 | 09/30/2026 | Open | Not discussed – carried forward. |" in md
    assert "| 2.01 | Order doors | GC | 09/08/2026 | 09/15/2026 | New |  |" in md
    assert "## Next Meeting\n09/15/2026, 9:00 AM, Trailer" in md
    assert firm.format.disclaimer in md


def test_docx_renders_and_reopens(tmp_path):
    firm, project, meeting, r2 = scenario()
    path = render_docx(firm, project, meeting, r2, tmp_path / "out" / "minutes.docx")
    assert path.exists()
    doc = Document(str(path))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "ACME OPM" in text and "OAC Meeting Minutes" in text
    assert "SAFETY" in text and "No incidents." in text
    assert "Approved PCO 1 at $100." in text
    assert "09/15/2026, 9:00 AM, Trailer" in text
    assert "Prepared by: Pat PM" in text
    tables = doc.tables
    assert len(tables) == 3  # header block, attendees, open items
    items_table = tables[-1]
    header = [c.text for c in items_table.rows[0].cells]
    assert header == ["Item", "Description", "Responsible", "Raised", "Due", "Status", "Latest update"]
    rows = [[c.text for c in r.cells] for r in items_table.rows[1:]]
    assert rows[0][0] == "1.01" and rows[0][5] == "OVERDUE"
    assert rows[1][5] == "Closed"
    assert "Mtg 1 (09/01/2026): Raised." in rows[0][1]  # history under the description
    assert "Mtg 2 (09/08/2026): Still late." in rows[0][1]
    assert doc.sections[0].footer.paragraphs[0].text.startswith("Town Library – Meeting No. 2 – 09/08/2026 – Page")


def test_docx_with_alternate_format(tmp_path):
    firm, project, meeting, r2 = scenario("action_log")
    path = render_docx(firm, project, meeting, r2, tmp_path / "log.docx")
    doc = Document(str(path))
    header = [c.text for c in doc.tables[-1].rows[0].cells]
    assert header == ["#", "Action", "Owner", "Due", "Status", "Update"]
    labels = [r.cells[4].text for r in doc.tables[-1].rows[1:]]
    assert labels == ["LATE", "Done", "Open", "New"]
    assert [r.cells[0].text for r in doc.tables[-1].rows[1:]] == ["A-001", "A-002", "A-003", "A-004"]


def test_xlsx_and_csv(tmp_path):
    firm, project, meeting, r2 = scenario()
    path = render_xlsx(r2.items, tmp_path / "log.xlsx", fmt=firm.format, as_of=date(2026, 9, 8), meeting_no=2, project_name="Town Library")
    wb = load_workbook(str(path))
    assert wb.sheetnames == ["Open Items", "All Items"]
    ws = wb["Open Items"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:7] == ("Item", "Description", "Section", "Responsible", "Raised", "Due", "Status")
    assert [r[0] for r in rows[1:]] == ["1.01", "1.03", "2.01"]
    assert rows[1][6] == "OVERDUE" and rows[1][7] == 3  # 3 days overdue on 9/8
    assert rows[1][2] == "Schedule"
    assert ws["E2"].number_format == "mm/dd/yyyy"
    assert ws.freeze_panes == "B2"
    all_rows = list(wb["All Items"].iter_rows(values_only=True))
    assert [r[0] for r in all_rows[1:]] == ["1.01", "1.02", "1.03", "2.01"]
    csv_text = render_csv(r2.items, fmt=firm.format, as_of=date(2026, 9, 8), meeting_no=2)
    assert csv_text.splitlines()[0].startswith("Item,Description,Section")
    assert "1.02,Answer RFI 4,RFIs & Submittals,Architect,2026-09-01,2026-09-20,Closed" in csv_text


def test_slug_and_basename():
    assert slugify("Bramford Fire Station Renovation & Addition") == "Bramford-Fire-Station-Renovation-Addition"
    assert output_basename("Town Library", 3, date(2026, 9, 15)) == "Town-Library_OAC-03_2026-09-15"
