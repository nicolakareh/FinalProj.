import io

from docx import Document

from oacminutes.parsers import kind_of, load_transcript, normalize_transcript, parse_srt, parse_vtt, word_count

ZOOM_VTT = """WEBVTT

1
00:00:01.000 --> 00:00:04.000
Dana Whitfield: Good morning everyone, let's get started.

2
00:00:04.500 --> 00:00:07.000
Dana Whitfield: First up is safety.

3
00:00:07.500 --> 00:00:10.000
Jorge Alvarez: No incidents this week.
"""

TEAMS_VTT = """WEBVTT

NOTE duration:"00:12:00.000"

a1b2c3d4-0001
00:00:01.000 --> 00:00:04.000
<v Dana Whitfield>Good morning everyone.</v>

a1b2c3d4-0002
00:00:04.500 --> 00:00:07.000
<v Jorge Alvarez>Morning. <i>No incidents</i> this week.</v>
"""

SRT = """1
00:00:01,000 --> 00:00:04,000
Dana Whitfield: Good morning everyone.

2
00:00:04,500 --> 00:00:07,000
Jorge Alvarez: Morning.
"""


def test_zoom_vtt_merges_same_speaker_runs():
    out = parse_vtt(ZOOM_VTT)
    assert out.splitlines() == [
        "Dana Whitfield: Good morning everyone, let's get started. First up is safety.",
        "Jorge Alvarez: No incidents this week.",
    ]


def test_teams_vtt_voice_tags_and_notes():
    out = parse_vtt(TEAMS_VTT)
    assert out.splitlines() == [
        "Dana Whitfield: Good morning everyone.",
        "Jorge Alvarez: Morning. No incidents this week.",
    ]


def test_srt():
    assert parse_srt(SRT).splitlines() == ["Dana Whitfield: Good morning everyone.", "Jorge Alvarez: Morning."]


def test_load_transcript_dispatches_on_extension():
    assert load_transcript(ZOOM_VTT.encode(), "meeting.vtt").startswith("Dana Whitfield:")
    assert load_transcript(SRT.encode(), "meeting.SRT").startswith("Dana Whitfield:")
    assert load_transcript(b"Line one\r\n\r\n\r\n\r\nLine two  \r\n", "notes.txt") == "Line one\n\nLine two"
    assert load_transcript("﻿BOM text".encode("utf-8"), "notes.md") == "BOM text"


def test_load_docx_transcript():
    doc = Document()
    doc.add_paragraph("Dana Whitfield: Welcome.")
    doc.add_paragraph("")
    doc.add_paragraph("Jorge Alvarez: Thanks.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Item"
    table.rows[0].cells[1].text = "Owner"
    buf = io.BytesIO()
    doc.save(buf)
    out = load_transcript(buf.getvalue(), "transcript.docx")
    assert "Dana Whitfield: Welcome." in out
    assert "Item | Owner" in out


def test_unsupported_extension():
    import pytest
    from oacminutes.parsers import UnsupportedFileType
    with pytest.raises(UnsupportedFileType):
        load_transcript(b"x", "file.xyz")


def test_kind_of():
    assert kind_of("a.vtt") == "transcript"
    assert kind_of("a.M4A") == "audio"
    assert kind_of("a.exe") == "unknown"


def test_normalize_and_word_count():
    assert normalize_transcript("a\r\nb\n\n\n\nc") == "a\nb\n\nc"
    assert word_count("one two  three\nfour") == 4
