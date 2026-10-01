import io
from types import SimpleNamespace

from docx import Document

from oacminutes.format_from_minutes import FormatProposal, HeaderProposal, SectionProposal, ColumnProposal, document_blocks, propose_format, to_firm_format
from oacminutes.models import FirmFormat


def proposal():
    return FormatProposal(
        name="Acme OPM", minutes_title="Owner's Meeting Minutes",
        header_fields=[HeaderProposal(key="project_name", label="Project"), HeaderProposal(key="meeting_number", label="Mtg No.")],
        sections=[SectionProposal(key="Safety & Site", title="Safety & Site", hint="incidents, logistics"), SectionProposal(key="cost", title="Cost", hint="PCOs, COs")],
        open_items_columns=[ColumnProposal(key="item_number", label="No."), ColumnProposal(key="description", label="Item"), ColumnProposal(key="responsible", label="Ball in court"), ColumnProposal(key="due_date", label="Required"), ColumnProposal(key="status", label="Status")],
        numbering_style="meeting_item", numbering_separator="-", numbering_pad=3,
        label_closed="Complete", label_overdue="LATE", disclaimer="Report corrections within 3 days.", date_format="%m.%d.%y", observations=["Ball-in-court column uses company names only."],
    )


def test_to_firm_format_fills_gaps_and_slugs_keys():
    fmt = to_firm_format(proposal())
    assert isinstance(fmt, FirmFormat)
    assert [s.key for s in fmt.sections] == ["safety_site", "cost", "general"]
    assert fmt.numbering.separator == "-" and fmt.numbering.item_pad == 3
    assert fmt.status_labels["closed"] == "Complete" and fmt.status_labels["overdue"] == "LATE"
    assert fmt.disclaimer == "Report corrections within 3 days."
    assert fmt.date_format == "%m.%d.%y"
    assert [c.label for c in fmt.open_items_columns][:3] == ["No.", "Item", "Ball in court"]


def test_document_blocks_docx_pdf_text():
    doc = Document(); doc.add_paragraph("OAC Meeting Minutes"); t = doc.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "1.01"; t.rows[0].cells[1].text = "Submit schedule"
    buf = io.BytesIO(); doc.save(buf)
    blocks = document_blocks(buf.getvalue(), "minutes.docx")
    assert blocks[0]["type"] == "text" and "1.01 | Submit schedule" in blocks[0]["text"]
    pdf = document_blocks(b"%PDF-1.4 fake", "minutes.pdf")
    assert pdf[0]["type"] == "document" and pdf[0]["source"]["media_type"] == "application/pdf"
    txt = document_blocks("Minutes\nSafety: none".encode(), "minutes.txt")
    assert "<minutes>" in txt[0]["text"]


def test_propose_format_with_fake_client():
    calls = []

    class Messages:
        def parse(self, **kwargs):
            calls.append(kwargs)
            return SimpleNamespace(stop_reason="end_turn", parsed_output=proposal(), content=[])

    client = SimpleNamespace(beta=SimpleNamespace(messages=Messages()))
    fmt, notes = propose_format(b"Minutes text", "minutes.txt", firm_name="Acme OPM", client=client)
    assert fmt.name == "Acme OPM" and notes == ["Ball-in-court column uses company names only."]
    assert calls[0]["output_format"] is FormatProposal and calls[0]["fallbacks"] == "default"
    assert calls[0]["messages"][0]["content"][0]["type"] == "text"
