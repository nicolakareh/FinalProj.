"""Turn uploaded transcript files into plain text the extractor can read.

Supported: .txt, .md, .vtt (Zoom, Teams, Google Meet), .srt, .docx.
Audio and video files are routed to :mod:`oacminutes.transcription`.
"""
from __future__ import annotations

import io
import re
from pathlib import Path

TRANSCRIPT_EXTENSIONS = {".txt", ".md", ".vtt", ".srt", ".docx"}
AUDIO_EXTENSIONS = {".mp3", ".m4a", ".wav", ".mp4", ".mov", ".aac", ".ogg", ".webm", ".flac", ".mpeg", ".mpga"}

_TIMESTAMP_LINE = re.compile(r"^\s*\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3}\s*-->\s*\d{1,2}:\d{2}(?::\d{2})?[.,]\d{1,3}")
_VOICE_TAG = re.compile(r"<v\s+([^>]+)>(.*?)(?:</v>|$)", re.DOTALL)
_HTML_TAG = re.compile(r"<[^>]+>")
_SPEAKER_LINE = re.compile(r"^([A-Z][\w .'\-]{0,40}?):\s+(.*)$")


class UnsupportedFileType(ValueError):
    pass


def kind_of(name: str) -> str:
    """'transcript', 'audio' or 'unknown' based on the file extension."""
    ext = Path(name).suffix.lower()
    if ext in TRANSCRIPT_EXTENSIONS:
        return "transcript"
    if ext in AUDIO_EXTENSIONS:
        return "audio"
    return "unknown"


def load_transcript(data: bytes, name: str) -> str:
    """Decode an uploaded transcript file into normalised plain text."""
    ext = Path(name).suffix.lower()
    if ext == ".docx":
        return normalize_transcript(_read_docx(data))
    text = _decode(data)
    if ext == ".vtt":
        return normalize_transcript(parse_vtt(text))
    if ext == ".srt":
        return normalize_transcript(parse_srt(text))
    if ext in (".txt", ".md"):
        return normalize_transcript(text)
    raise UnsupportedFileType(f"Unsupported transcript type: {ext or name}")


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _read_docx(data: bytes) -> str:
    from docx import Document  # python-docx

    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def parse_vtt(text: str) -> str:
    """WebVTT -> 'Speaker: text' lines. Handles Zoom, Teams (<v> tags) and Meet."""
    lines: list[str] = []
    in_note = False
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            in_note = False
            continue
        if line.startswith("WEBVTT") or line.startswith("NOTE") or line.startswith("STYLE") or line.startswith("REGION"):
            in_note = True
            continue
        if in_note:
            continue
        if _TIMESTAMP_LINE.match(line):
            continue
        if re.fullmatch(r"\d+", line) or re.fullmatch(r"[0-9a-f-]{8,}", line):
            continue  # cue identifier
        voices = _VOICE_TAG.findall(line)
        if voices:
            for speaker, spoken in voices:
                spoken = _HTML_TAG.sub("", spoken).strip()
                if spoken:
                    lines.append(f"{speaker.strip()}: {spoken}")
            continue
        lines.append(_HTML_TAG.sub("", line))
    return _merge_speaker_runs(lines)


def parse_srt(text: str) -> str:
    lines: list[str] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line or re.fullmatch(r"\d+", line) or _TIMESTAMP_LINE.match(line):
            continue
        lines.append(_HTML_TAG.sub("", line))
    return _merge_speaker_runs(lines)


def _merge_speaker_runs(lines: list[str]) -> str:
    """Join consecutive cues from the same speaker into one paragraph."""
    merged: list[tuple[str | None, str]] = []
    for line in lines:
        m = _SPEAKER_LINE.match(line)
        speaker, spoken = (m.group(1).strip(), m.group(2).strip()) if m else (None, line)
        if merged and merged[-1][0] == speaker and speaker is not None:
            merged[-1] = (speaker, merged[-1][1] + " " + spoken)
        elif merged and merged[-1][0] is None and speaker is None:
            merged[-1] = (None, merged[-1][1] + " " + spoken)
        else:
            merged.append((speaker, spoken))
    return "\n".join(f"{s}: {t}" if s else t for s, t in merged)


def normalize_transcript(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def word_count(text: str) -> int:
    return len(re.findall(r"\S+", text))
