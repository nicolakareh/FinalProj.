"""Propose a firm format from one of the firm's own issued minutes documents.

This is how a new OPM firm gets onboarded: upload a sample of their minutes, review the
proposed format file, save it. Uses Claude with structured output; the result is a
`FirmFormat` the rest of the app already understands.

Command line:  python -m oacminutes.format_from_minutes minutes.docx --out data/formats/firm.json
"""
from __future__ import annotations

import argparse
import base64
import io
import json
import sys
from pathlib import Path
from typing import Literal, Optional

from pydantic import BaseModel, Field

from .extract import DEFAULT_MODEL, ExtractionError
from .models import ColumnSpec, FirmFormat, HeaderField, NumberingSpec, SectionSpec

HEADER_KEYS = Literal["project_name", "project_number", "meeting_number", "meeting_date", "meeting_time", "location", "owner", "architect", "contractor", "opm_firm", "prepared_by", "distribution"]
COLUMN_KEYS = Literal["item_number", "description", "section", "responsible", "date_raised", "due_date", "status", "notes", "last_update"]


class HeaderProposal(BaseModel):
    key: HEADER_KEYS
    label: str = Field(description="The label exactly as printed in the minutes")


class SectionProposal(BaseModel):
    key: str = Field(description="snake_case identifier")
    title: str = Field(description="Section heading exactly as printed")
    hint: str = Field(description="What the firm files under this heading, in the firm's own vocabulary; one line")


class ColumnProposal(BaseModel):
    key: COLUMN_KEYS
    label: str = Field(description="Column header exactly as printed")


class FormatProposal(BaseModel):
    name: str = Field(description="Short name for the format, usually the firm's name")
    minutes_title: str = Field(description="The document title as printed, e.g. 'OAC Meeting Minutes'")
    header_fields: list[HeaderProposal] = Field(description="The header block fields in the order printed")
    sections: list[SectionProposal] = Field(description="Discussion sections in the order they appear; omit the open-items table itself")
    open_items_columns: list[ColumnProposal] = Field(description="Columns of the open/action items table in order")
    numbering_style: Literal["meeting_item", "sequential"] = Field(description="meeting_item when items read like 12.03 (meeting.sequence); sequential when they are a running count")
    numbering_prefix: str = Field(default="", description="Any prefix before item numbers, e.g. 'A-'")
    numbering_separator: str = Field(default=".", description="Separator between meeting and sequence for meeting_item style")
    numbering_pad: int = Field(default=2, description="Digits used for the sequence part, e.g. 2 for 12.03")
    label_open: str = "Open"
    label_on_hold: str = "On Hold"
    label_closed: str = "Closed"
    label_overdue: str = "OVERDUE"
    label_new: str = "New"
    not_discussed_text: str = Field(default="Not discussed – carried forward.", description="How the firm marks an item that was not discussed")
    disclaimer: str = Field(default="", description="Closing disclaimer or correction-period language, verbatim; empty if none")
    date_format: str = Field(default="%m/%d/%Y", description="Python strftime pattern matching how dates are printed")
    font_name: str = Field(default="Calibri", description="Body typeface if identifiable")
    show_attendees_table: bool = True
    show_decisions: bool = Field(default=True, description="Whether the minutes have a separate decisions list")
    observations: list[str] = Field(default_factory=list, description="Anything about the house style the schema cannot capture, for the person reviewing this proposal")


def to_firm_format(p: FormatProposal) -> FirmFormat:
    sections = [SectionSpec(key=_slug(s.key or s.title), title=s.title.strip(), hint=s.hint.strip()) for s in p.sections if s.title.strip()]
    if not any(s.key == "general" for s in sections):
        sections.append(SectionSpec(key="general", title="General", hint="anything that does not fit above"))
    columns = [ColumnSpec(key=c.key, label=c.label.strip() or c.key) for c in p.open_items_columns] or FirmFormat().open_items_columns
    return FirmFormat(
        name=p.name.strip() or "Firm format",
        minutes_title=p.minutes_title.strip() or "Meeting Minutes",
        header_fields=[HeaderField(key=h.key, label=h.label.strip() or h.key) for h in p.header_fields],
        sections=sections,
        open_items_columns=columns,
        numbering=NumberingSpec(style=p.numbering_style, prefix=p.numbering_prefix, separator=p.numbering_separator or ".", item_pad=max(1, min(4, p.numbering_pad))),
        status_labels={"open": p.label_open, "on_hold": p.label_on_hold, "closed": p.label_closed, "overdue": p.label_overdue, "new": p.label_new},
        show_attendees_table=p.show_attendees_table,
        show_decisions=p.show_decisions,
        not_discussed_text=p.not_discussed_text or "Not discussed – carried forward.",
        disclaimer=p.disclaimer.strip() or FirmFormat().disclaimer,
        date_format=p.date_format if "%" in p.date_format else "%m/%d/%Y",
        font_name=p.font_name.strip() or "Calibri",
    )


def _slug(text: str) -> str:
    import re
    slug = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    return slug or "section"


SYSTEM = """You are given one issued construction meeting minutes document from an Owner's Project Manager firm. Describe the firm's house format precisely enough that software can reproduce it for future meetings. Copy labels, headings and column names exactly as printed. Section hints must say what topics the firm files under each heading, in the firm's own words. Infer the item numbering style from the item numbers you see. Quote the disclaimer verbatim. Put anything the schema cannot express in observations."""


def document_blocks(data: bytes, name: str) -> list[dict]:
    """Turn a minutes file into message content: PDFs go as documents, everything else as text."""
    ext = Path(name).suffix.lower()
    if ext == ".pdf":
        return [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": base64.standard_b64encode(data).decode("ascii")}}]
    if ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    lines.append(" | ".join(cells))
        text = "\n".join(lines)
    else:
        text = data.decode("utf-8", errors="replace")
    if not text.strip():
        raise ExtractionError("The sample document has no readable text.")
    return [{"type": "text", "text": "<minutes>\n" + text.strip() + "\n</minutes>"}]


def propose_format(data: bytes, name: str, *, firm_name: str = "", client=None, model: str = DEFAULT_MODEL, use_fallbacks: bool = True) -> tuple[FirmFormat, list[str]]:
    """Return (format, observations) for a sample minutes document."""
    import anthropic

    client = client or anthropic.Anthropic()
    content = document_blocks(data, name)
    content.append({"type": "text", "text": "Describe this firm's minutes format." + (f" The firm is {firm_name}." if firm_name else "")})
    kwargs: dict = dict(model=model, max_tokens=8000, system=SYSTEM, messages=[{"role": "user", "content": content}], output_format=FormatProposal, output_config={"effort": "high"})
    if use_fallbacks:
        kwargs["betas"] = ["server-side-fallback-2026-07-01"]
        kwargs["fallbacks"] = "default"
    try:
        message = client.beta.messages.parse(**kwargs)
    except anthropic.AuthenticationError as exc:
        raise ExtractionError("Claude API key is missing or invalid. Set ANTHROPIC_API_KEY and restart.") from exc
    except anthropic.APIStatusError as exc:
        raise ExtractionError(f"Claude API error {exc.status_code}: {exc.message}") from exc
    except anthropic.APIConnectionError as exc:
        raise ExtractionError("Could not reach the Claude API. Check the network connection.") from exc
    if message.stop_reason == "refusal":
        raise ExtractionError("The model declined to read this document.")
    proposal = getattr(message, "parsed_output", None)
    if proposal is None:
        text = next((b.text for b in message.content if getattr(b, "type", "") == "text"), "")
        proposal = FormatProposal.model_validate(json.loads(text))
    fmt = to_firm_format(proposal)
    if firm_name and (not proposal.name.strip() or proposal.name.strip().lower() in ("firm format", "format")):
        fmt.name = firm_name
    return fmt, list(proposal.observations)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Propose a firm format from a sample minutes document.")
    parser.add_argument("sample", help="Issued minutes as .docx, .pdf, .txt or .md")
    parser.add_argument("--out", required=True, help="Where to write the format JSON")
    parser.add_argument("--firm", default="", help="Firm name")
    args = parser.parse_args(argv)
    data = Path(args.sample).read_bytes()
    fmt, observations = propose_format(data, Path(args.sample).name, firm_name=args.firm)
    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(fmt.model_dump_json(indent=2), encoding="utf-8")
    print(f"Wrote {out} ({len(fmt.sections)} sections, {len(fmt.open_items_columns)} columns, numbering {fmt.numbering.style}).")
    for note in observations:
        print(" -", note)
    return 0


if __name__ == "__main__":
    sys.exit(main())
