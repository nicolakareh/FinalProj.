"""Word rendering of the minutes in the firm's format."""
from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .carryforward import CarryForwardResult
from .dates import fmt_date, parse_date
from .models import Firm, FirmFormat, Meeting, Project
from .render_common import cell_text, history_lines, row_kind, sections_with_content, visible_items

GRAY = RGBColor(0x80, 0x80, 0x80)
RED = RGBColor(0xC0, 0x00, 0x00)


def _shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def _set_col_widths(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)


def _add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _heading(doc, text: str, accent: RGBColor, size: int) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(size)
    run.font.color.rgb = accent


def _fill_cell(cell, text: str, *, bold: bool = False, color: RGBColor | None = None, size: int | None = None, italic: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    if size is not None:
        run.font.size = Pt(size)


def render_docx(firm: Firm, project: Project, meeting: Meeting, result: CarryForwardResult, path: str | Path) -> Path:
    fmt: FirmFormat = firm.format
    accent = RGBColor.from_string(fmt.accent_color.upper())
    doc = Document()

    normal = doc.styles["Normal"]
    normal.font.name = fmt.font_name
    normal.font.size = Pt(fmt.font_size)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), fmt.font_name)

    section = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Inches(0.75))
    usable_in = (section.page_width - section.left_margin - section.right_margin) / 914400

    # ---- title block
    p = doc.add_paragraph()
    run = p.add_run(firm.name.upper())
    run.bold = True
    run.font.size = Pt(fmt.font_size + 5)
    run.font.color.rgb = accent
    p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    run = p.add_run(fmt.minutes_title)
    run.bold = True
    run.font.size = Pt(fmt.font_size + 3)

    fields = [(hf.label, project.header_value(hf.key, meeting)) for hf in fmt.header_fields]
    if fields:
        rows = math.ceil(len(fields) / 2)
        table = doc.add_table(rows=rows, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.LEFT
        for idx, (label, value) in enumerate(fields):
            row, col = divmod(idx, 2)
            _fill_cell(table.cell(row, col * 2), f"{label}:", bold=True)
            _fill_cell(table.cell(row, col * 2 + 1), value)
        quarter = usable_in / 4
        _set_col_widths(table, [quarter * 0.7, quarter * 1.3, quarter * 0.7, quarter * 1.3])

    extraction = meeting.extraction

    # ---- attendees
    if fmt.show_attendees_table and extraction and extraction.attendees:
        _heading(doc, "Attendees", accent, fmt.font_size + 1)
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        for cell, label in zip(table.rows[0].cells, ["Name", "Company", "Role", "Present"]):
            _fill_cell(cell, label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade(cell, fmt.accent_color)
        for a in extraction.attendees:
            cells = table.add_row().cells
            _fill_cell(cells[0], a.name)
            _fill_cell(cells[1], a.company)
            _fill_cell(cells[2], a.role)
            _fill_cell(cells[3], "Yes" if a.present else "No")
        _set_col_widths(table, [usable_in * 0.3, usable_in * 0.3, usable_in * 0.28, usable_in * 0.12])

    # ---- discussion sections
    for title, bullets in sections_with_content(extraction, fmt):
        _heading(doc, title, accent, fmt.font_size + 1)
        for bullet in bullets:
            para = doc.add_paragraph(bullet, style="List Bullet")
            para.paragraph_format.space_after = Pt(2)

    if fmt.show_decisions and extraction and extraction.decisions:
        _heading(doc, "Decisions", accent, fmt.font_size + 1)
        for decision in extraction.decisions:
            para = doc.add_paragraph(decision, style="List Number")
            para.paragraph_format.space_after = Pt(2)

    # ---- open items
    _heading(doc, "Open Items", accent, fmt.font_size + 1)
    items = visible_items(result, fmt)
    if not items:
        doc.add_paragraph("No open items.")
    else:
        cols = fmt.open_items_columns
        table = doc.add_table(rows=1, cols=len(cols))
        table.style = "Table Grid"
        for cell, col in zip(table.rows[0].cells, cols):
            _fill_cell(cell, col.label, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
            _shade(cell, fmt.accent_color)
        has_notes_col = any(c.key == "notes" for c in cols)
        for item in items:
            kind = row_kind(item, meeting_no=meeting.meeting_no, as_of=meeting.meeting_date)
            cells = table.add_row().cells
            for cell, col in zip(cells, cols):
                text = cell_text(item, col.key, fmt, meeting_no=meeting.meeting_no, as_of=meeting.meeting_date)
                color = GRAY if kind == "closed" else None
                bold = False
                if col.key == "status" and kind == "overdue":
                    color, bold = RED, True
                if col.key == "due_date" and kind == "overdue":
                    color = RED
                _fill_cell(cell, text, bold=bold, color=color)
                if col.key == "description" and fmt.show_item_history and not has_notes_col and len(item.history) > 1:
                    for line in history_lines(item, fmt):
                        para = cell.add_paragraph()
                        para.paragraph_format.space_after = Pt(0)
                        run = para.add_run(line)
                        run.font.size = Pt(max(fmt.font_size - 2, 7))
                        run.font.color.rgb = GRAY if kind == "closed" else RGBColor(0x40, 0x40, 0x40)
            if kind == "closed":
                for cell in cells:
                    _shade(cell, "F2F2F2")
            elif kind == "overdue":
                for cell in cells:
                    _shade(cell, "FDECEC")
        widths = [c.width_in or 1.0 for c in cols]
        scale = usable_in / sum(widths)
        _set_col_widths(table, [w * scale for w in widths])
        legend = doc.add_paragraph()
        legend.paragraph_format.space_before = Pt(3)
        run = legend.add_run(
            f"Status legend: {fmt.status_labels.get('new', 'New')} = raised this meeting; "
            f"{fmt.status_labels.get('overdue', 'OVERDUE')} = past due; "
            f"{fmt.status_labels.get('closed', 'Closed')} items appear once and are then removed from the log."
        )
        run.italic = True
        run.font.size = Pt(max(fmt.font_size - 2, 7))

    # ---- next meeting
    if extraction and extraction.next_meeting:
        nm = extraction.next_meeting
        when = fmt_date(parse_date(nm.date, meeting.meeting_date.year), fmt.date_format) or (nm.date or "")
        parts = [p for p in [when, nm.time, nm.location] if p]
        if parts:
            _heading(doc, "Next Meeting", accent, fmt.font_size + 1)
            doc.add_paragraph(", ".join(parts))

    # ---- disclaimer
    para = doc.add_paragraph()
    para.paragraph_format.space_before = Pt(12)
    run = para.add_run(fmt.disclaimer)
    run.italic = True
    run.font.size = Pt(max(fmt.font_size - 1, 7))
    if project.prepared_by:
        para = doc.add_paragraph()
        run = para.add_run(f"Prepared by: {project.prepared_by}")
        run.font.size = Pt(max(fmt.font_size - 1, 7))

    # ---- footer
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{project.name} – Meeting No. {meeting.meeting_no} – {fmt_date(meeting.meeting_date, fmt.date_format)} – Page ")
    run.font.size = Pt(8)
    _add_page_number(footer)

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path
