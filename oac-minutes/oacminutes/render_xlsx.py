"""Open-items log as an Excel workbook (and CSV)."""
from __future__ import annotations

import csv
import io
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .carryforward import parse_item_number
from .models import FirmFormat, ItemStatus, OpenItem
from .render_common import history_lines, status_label

COLUMNS = [
    ("Item", 9), ("Description", 48), ("Section", 22), ("Responsible", 22), ("Raised", 12),
    ("Due", 12), ("Status", 12), ("Days overdue", 12), ("Meetings open", 13),
    ("Last discussed (Mtg)", 12), ("Latest update", 48), ("History", 60),
]


def _excel_date_format(strftime_pattern: str) -> str:
    return (
        strftime_pattern.replace("%m", "mm").replace("%d", "dd").replace("%Y", "yyyy").replace("%y", "yy")
    )


def _rows(items: list[OpenItem], fmt: FirmFormat, *, as_of: date, meeting_no: int) -> list[list]:
    rows = []
    for item in sorted(items, key=lambda i: parse_item_number(i.item_number)):
        overdue = item.is_overdue(as_of)
        rows.append([
            item.item_number,
            item.description,
            fmt.section_title(item.section_key),
            item.responsible,
            item.date_raised,
            item.due_date,
            status_label(item, fmt, meeting_no=meeting_no, as_of=as_of),
            (as_of - item.due_date).days if overdue and item.due_date else 0,
            item.meetings_open(meeting_no),
            item.last_discussed_meeting_no,
            item.latest_note(),
            "\n".join(history_lines(item, fmt)),
        ])
    return rows


def _write_sheet(ws, rows: list[list], fmt: FirmFormat) -> None:
    header_fill = PatternFill("solid", fgColor=fmt.accent_color)
    overdue_fill = PatternFill("solid", fgColor="FDECEC")
    closed_fill = PatternFill("solid", fgColor="F2F2F2")
    date_fmt = _excel_date_format(fmt.date_format)
    ws.append([c[0] for c in COLUMNS])
    for idx, (_, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in rows:
        ws.append(row)
        r = ws.max_row
        for col in (5, 6):
            ws.cell(row=r, column=col).number_format = date_fmt
        for col in (2, 11, 12):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        for col in (1, 3, 4, 5, 6, 7, 8, 9, 10):
            ws.cell(row=r, column=col).alignment = Alignment(vertical="top")
        status = row[6]
        if status == fmt.status_labels.get("overdue", "OVERDUE"):
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=col).fill = overdue_fill
            ws.cell(row=r, column=7).font = Font(bold=True, color="C00000")
        elif status == fmt.status_labels.get("closed", "Closed"):
            for col in range(1, len(COLUMNS) + 1):
                ws.cell(row=r, column=col).fill = closed_fill
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{max(ws.max_row, 1)}"


def render_xlsx(items: list[OpenItem], path: str | Path, *, fmt: FirmFormat, as_of: date, meeting_no: int, project_name: str = "") -> Path:
    wb = Workbook()
    ws_open = wb.active
    ws_open.title = "Open Items"
    open_items = [i for i in items if i.status != ItemStatus.CLOSED]
    _write_sheet(ws_open, _rows(open_items, fmt, as_of=as_of, meeting_no=meeting_no), fmt)
    ws_all = wb.create_sheet("All Items")
    _write_sheet(ws_all, _rows(list(items), fmt, as_of=as_of, meeting_no=meeting_no), fmt)
    wb.properties.title = f"{project_name} – Open Items Log".strip(" –")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return path


def render_csv(items: list[OpenItem], *, fmt: FirmFormat, as_of: date, meeting_no: int) -> str:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([c[0] for c in COLUMNS])
    for row in _rows(list(items), fmt, as_of=as_of, meeting_no=meeting_no):
        writer.writerow([v.isoformat() if isinstance(v, date) else v for v in row])
    return buf.getvalue()
